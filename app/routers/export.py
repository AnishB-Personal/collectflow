"""
Authoritative PostgreSQL Data Export Router for CollectFlow.

Provides:
- POST /api/export/token — Short-lived single-use export ticket issuance (admin, finance only).
- GET /api/export/excel — 17-column openpyxl Excel (.xlsx) workbook streaming.
- GET /api/export/csv — 17-column sanitized CSV streaming.

Security & Integrity:
- Stored ticket SHA-256 hashes, 5-minute TTL.
- Atomic single-use ticket consumption under PostgreSQL row lock (SELECT FOR UPDATE).
- Formula injection protection (' prefix on =, +, -, @).
- Exact 17-column order matching frontend AUTHORITATIVE_COLUMNS specification.
- Full reconciliation between Excel, CSV, and reporting results.
"""

import csv
import datetime
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import io
import secrets
from typing import Any, Dict, List, Optional
from fastapi import APIRouter, Depends, Header, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, or_, select

from app.core.database import get_db_session
from app.models.domain_models import (
    AuditLog,
    CommercialInvoice,
    Customer,
    ExportTicket,
    PaymentTransaction,
    PurchaseOrder,
    User,
)
from app.routers.auth import get_current_user, require_finance_or_admin

router = APIRouter(prefix="/api/export", tags=["Export"])

EXCEL_HEADER_FILL = "1F4E79"
EXCEL_CURRENCY_FORMAT = '"₹"#,##0.00'


class ExportTokenRequest(BaseModel):
    format: Optional[str] = None
    fy: Optional[str] = None
    scope: Optional[str] = None

AUTHORITATIVE_COLUMNS = [
    "FY",
    "PO Date",
    "PO Number",
    "Customer Name",
    "Executive",
    "Exclusive Amount (₹)",
    "Tax Amount (₹)",
    "Total Invoiced (₹)",
    "Invoice Number",
    "Invoice Date",
    "Due Date",
    "Cash Received (₹)",
    "TDS Deducted (₹)",
    "Total Settled (₹)",
    "Balance Outstanding (₹)",
    "Status",
    "Days Overdue",
]


def _quantize_2dec(val: Any) -> Decimal:
    if val is None:
        return Decimal("0.00")
    return Decimal(str(val)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _format_money(val: Any) -> str:
    if val is None:
        return "0.00"
    return f"{Decimal(str(val)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP):.2f}"


def sanitize_formula_injection(val: Any) -> Any:
    """Prepend single quote to neutralize formula injection in spreadsheet applications."""
    if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
        return "'" + val
    return val


@router.post("/token")
def create_export_token(
    request: Request,
    payload: Optional[ExportTokenRequest] = None,
    user: dict = Depends(require_finance_or_admin),
):
    """
    Generate a short-lived, single-use export ticket:
    - Admin and finance only.
    - 256 bits CSPRNG entropy.
    - Strictly bound to format, FY, and scope.
    - Stored as SHA-256 hash in export_tickets table (raw secret never stored or logged).
    - Expires in 5 minutes (300 seconds).
    - Atomically consumed on download under PostgreSQL row lock.
    """
    if not payload:
        raise HTTPException(
            status_code=400,
            detail="Export ticket request requires payload with explicit 'format', 'fy', and 'scope'.",
        )

    # 1. Format
    if not payload.format or not payload.format.strip():
        req_format = "excel"
    else:
        req_format = payload.format.lower().strip()
    if req_format not in ("excel", "csv"):
        raise HTTPException(status_code=400, detail="Invalid export format. Allowed formats: 'excel', 'csv'.")

    # 2. FY
    if not payload.fy or not payload.fy.strip():
        raise HTTPException(
            status_code=400,
            detail="Export ticket requires an explicit financial year (e.g. 'FY 2026-27') or explicit 'ALL'. Unscoped tickets are prohibited.",
        )
    raw_fy = payload.fy.strip()
    req_fy = "ALL" if raw_fy.upper() in ("ALL", "ALL FY", "ALL FINANCIAL YEARS") else raw_fy

    # 3. Scope
    if not payload.scope or not payload.scope.strip():
        raise HTTPException(
            status_code=400,
            detail="Export ticket requires an explicit scope (e.g. 'all', 'outstanding'). Unscoped tickets are prohibited.",
        )
    raw_scope = payload.scope.strip().lower()
    req_scope = "all" if raw_scope in ("all", "all_records", "all records", "all invoices") else raw_scope

    ticket_secret = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(ticket_secret.encode("utf-8")).hexdigest()
    now_utc = datetime.datetime.now(datetime.timezone.utc)
    expires_at = now_utc + datetime.timedelta(minutes=5)

    with get_db_session() as db:
        ticket_record = ExportTicket(
            token_hash=token_hash,
            user_id=user["id"],
            created_at=now_utc,
            expires_at=expires_at,
            used=False,
            format=req_format,
            fy_code=req_fy,
            scope=req_scope,
        )
        db.add(ticket_record)

        audit_entry = AuditLog(
            user_id=user["id"],
            username=user["username"],
            action="EXPORT_TOKEN_ISSUED",
            entity_type="export_ticket",
            entity_id=token_hash[:16],
            details={
                "expires_in": 300,
                "issued_to": user["username"],
                "role": user["role"],
                "format": req_format,
                "fy": req_fy,
                "scope": req_scope,
            },
            ip_address=request.client.host if request.client else None,
            request_id=getattr(request.state, "request_id", None),
        )
        db.add(audit_entry)
        db.commit()

    return {
        "ticket": ticket_secret,
        "token": ticket_secret,
        "expires_in": 300,
        "expires_at": expires_at.isoformat(),
        "format": req_format,
        "fy": req_fy,
        "scope": req_scope,
    }


def authenticate_and_consume_export_ticket(
    request: Request,
    ticket: Optional[str] = Query(None),
    authorization: Optional[str] = Header(None),
) -> dict:
    """
    Authenticates export requests:
    - Every download requires a valid single-use export ticket; bearer authorization alone
      CANNOT bypass ticket validation.
    - Every download requires a valid single-use export ticket.
    - User identity and active session required (caller authenticated via session cookie or bearer).
    - Caller must have admin or finance role and must not have pending password rotation.
    - Caller identity must match the user to whom the ticket was issued.
    - Format, FY and scope validated against exact canonical values recorded at issuance.
    - Atomically consumed under PostgreSQL row lock (SELECT FOR UPDATE) ONLY after all authorization checks pass.
    """
    if not ticket or not ticket.strip():
        raise HTTPException(
            status_code=401,
            detail="Export download requires a valid single-use export ticket. Bearer authorization alone cannot bypass ticket requirement.",
        )

    # 1. Authenticate caller (cookie session or bearer) and enforce active status + password rotation policy
    caller = get_current_user(request=request, authorization=authorization)
    if caller.get("role") not in ("admin", "finance"):
        raise HTTPException(
            status_code=403,
            detail="Export generation requires admin or finance privileges.",
        )

    req_path = request.url.path.lower()
    req_format = "excel" if "excel" in req_path else "csv" if "csv" in req_path else None
    req_fy = request.query_params.get("fy")
    req_scope = request.query_params.get("scope")

    token_hash = hashlib.sha256(ticket.strip().encode("utf-8")).hexdigest()
    now = datetime.datetime.now(datetime.timezone.utc)

    with get_db_session() as db:
        ticket_row = (
            db.query(ExportTicket)
            .filter(ExportTicket.token_hash == token_hash)
            .with_for_update()
            .first()
        )

        if not ticket_row:
            raise HTTPException(
                status_code=401,
                detail="Invalid, expired, or already consumed export ticket.",
            )

        if ticket_row.used:
            raise HTTPException(
                status_code=401,
                detail="Invalid, expired, or already consumed export ticket.",
            )

        exp_time = ticket_row.expires_at
        if exp_time.tzinfo is None:
            exp_time = exp_time.replace(tzinfo=datetime.timezone.utc)

        if exp_time < now:
            raise HTTPException(
                status_code=401,
                detail="Invalid, expired, or already consumed export ticket.",
            )

        # 2. Verify ticket ownership: caller must match ticket recipient
        if caller.get("id") != ticket_row.user_id:
            raise HTTPException(
                status_code=403,
                detail="Export ticket was issued to a different user.",
            )

        owner = db.query(User).filter(User.id == ticket_row.user_id).first()
        if not owner or not owner.is_active or owner.role not in ("admin", "finance"):
            raise HTTPException(
                status_code=403,
                detail="Export generation requires admin or finance privileges.",
            )

        # Format binding check
        if ticket_row.format and req_format:
            if ticket_row.format.lower() != req_format.lower():
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to format '{ticket_row.format}', cannot download as '{req_format}'.",
                )

        # FY binding check against exact canonical values
        raw_dl_fy = req_fy.strip() if req_fy and req_fy.strip() else None
        dl_fy_canon = ("ALL" if raw_dl_fy.upper() in ("ALL", "ALL FY") else raw_dl_fy) if raw_dl_fy else None

        if ticket_row.fy_code == "ALL":
            if dl_fy_canon != "ALL":
                raise HTTPException(
                    status_code=400,
                    detail="Export ticket was issued for financial year 'ALL'. Request must specify fy='ALL'.",
                )
        else:
            if not dl_fy_canon:
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to financial year '{ticket_row.fy_code}'. Missing 'fy' parameter.",
                )
            if dl_fy_canon == "ALL":
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to financial year '{ticket_row.fy_code}', cannot request scope 'ALL'.",
                )
            if dl_fy_canon.upper() != ticket_row.fy_code.upper():
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to financial year '{ticket_row.fy_code}', cannot request '{dl_fy_canon}'.",
                )

        # Scope binding check against exact canonical values
        raw_dl_scope = req_scope.strip().lower() if req_scope and req_scope.strip() else None
        dl_scope_canon = ("all" if raw_dl_scope in ("all", "all_records", "all records") else raw_dl_scope) if raw_dl_scope else None

        if ticket_row.scope == "all":
            if dl_scope_canon != "all":
                raise HTTPException(
                    status_code=400,
                    detail="Export ticket was issued for scope 'all'. Request must specify scope='all'.",
                )
        else:
            if not dl_scope_canon:
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to scope '{ticket_row.scope}'. Missing 'scope' parameter.",
                )
            if dl_scope_canon in ("all", "all_records", "all records"):
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to scope '{ticket_row.scope}', cannot request 'all'.",
                )
            if dl_scope_canon != ticket_row.scope.lower():
                raise HTTPException(
                    status_code=400,
                    detail=f"Export ticket is bound to scope '{ticket_row.scope}', cannot request '{dl_scope_canon}'.",
                )

        # Atomically mark as consumed under row lock
        ticket_row.used = True

        audit_entry = AuditLog(
            user_id=owner.id,
            username=owner.username,
            action="EXPORT_DOWNLOAD_COMPLETED",
            entity_type="export_ticket",
            entity_id=token_hash[:16],
            details={
                "download_by": owner.username,
                "role": owner.role,
                "format": req_format,
                "fy": dl_fy_canon,
                "scope": dl_scope_canon,
            },
            ip_address=request.client.host if request.client else None,
            request_id=getattr(request.state, "request_id", None),
        )
        db.add(audit_entry)
        db.commit()

        return {
            "id": owner.id,
            "username": owner.username,
            "role": owner.role,
            "full_name": owner.full_name,
        }


def _compile_export_records(
    fy: str = "ALL",
    scope: str = "all",
    as_of_date: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Compiles records into 17 authoritative columns matching PostgreSQL state."""
    ref_date = (
        datetime.date.fromisoformat(as_of_date)
        if as_of_date
        else datetime.date.today()
    )

    records: List[Dict[str, Any]] = []

    with get_db_session() as db:
        # Preload settlements as of ref_date
        pay_stmt = select(PaymentTransaction)
        if as_of_date:
            pay_stmt = pay_stmt.where(PaymentTransaction.payment_date <= ref_date)
        all_payments = db.execute(pay_stmt).scalars().all()
        invoice_settlements: Dict[int, Dict[str, Decimal]] = {}
        for p in all_payments:
            inv_id = p.invoice_id
            if inv_id not in invoice_settlements:
                invoice_settlements[inv_id] = {
                    "cash": Decimal("0.00"),
                    "tds": Decimal("0.00"),
                    "total": Decimal("0.00"),
                }
            invoice_settlements[inv_id]["cash"] += _quantize_2dec(p.cash_amount)
            invoice_settlements[inv_id]["tds"] += _quantize_2dec(p.tds_amount)
            invoice_settlements[inv_id]["total"] += _quantize_2dec(p.total_settled)

        # 1. Commercial Invoices
        inv_stmt = (
            select(CommercialInvoice, Customer, User, PurchaseOrder)
            .join(Customer, CommercialInvoice.customer_id == Customer.id)
            .join(User, CommercialInvoice.executive_id == User.id)
            .outerjoin(PurchaseOrder, CommercialInvoice.po_id == PurchaseOrder.id)
            .order_by(CommercialInvoice.id.desc())
        )
        if fy and fy.upper() not in ("ALL", "ALL FY"):
            inv_stmt = inv_stmt.where(CommercialInvoice.fy_code == fy)

        if as_of_date:
            inv_stmt = inv_stmt.where(
                CommercialInvoice.invoice_date <= ref_date,
                or_(
                    CommercialInvoice.lifecycle_status == "ISSUED",
                    (CommercialInvoice.lifecycle_status == "CANCELLED")
                    & (func.date(CommercialInvoice.cancelled_at) > ref_date),
                ),
            )

        inv_rows = db.execute(inv_stmt).all()

        for inv, cust, exec_user, po in inv_rows:
            st = invoice_settlements.get(
                inv.id,
                {"cash": Decimal("0.00"), "tds": Decimal("0.00"), "total": Decimal("0.00")},
            )

            # Cancellation status as of ref_date
            if as_of_date and inv.lifecycle_status == "CANCELLED":
                canc_dt = inv.cancelled_at.date() if inv.cancelled_at else None
                is_cancelled = canc_dt is not None and canc_dt <= ref_date
            else:
                is_cancelled = inv.lifecycle_status == "CANCELLED"

            tot_amt = Decimal("0.00") if is_cancelled else _quantize_2dec(inv.total_amount)
            excl_amt = Decimal("0.00") if is_cancelled else _quantize_2dec(inv.exclusive_amount)
            tax_amt = Decimal("0.00") if is_cancelled else _quantize_2dec(inv.tax_amount)
            bal = Decimal("0.00") if is_cancelled else max(Decimal("0.00"), tot_amt - st["total"])

            if scope == "outstanding" and bal <= Decimal("0.00"):
                continue

            days_overdue = 0
            if not is_cancelled and bal > Decimal("0.00") and inv.due_date < ref_date:
                days_overdue = max(0, (ref_date - inv.due_date).days)

            status_str = (
                "CANCELLED"
                if is_cancelled
                else "PAID"
                if bal == Decimal("0.00")
                else "PARTIAL"
                if st["total"] > Decimal("0.00")
                else "PENDING"
            )

            po_number = po.po_number if (po and po.po_number) else "—"
            po_date_str = str(po.po_date) if (po and po.po_date) else str(inv.invoice_date)

            records.append(
                {
                    "FY": sanitize_formula_injection(inv.fy_code),
                    "PO Date": str(po_date_str),
                    "PO Number": sanitize_formula_injection(po_number),
                    "Customer Name": sanitize_formula_injection(cust.billing_name),
                    "Executive": sanitize_formula_injection(exec_user.full_name or exec_user.username),
                    "Exclusive Amount (₹)": float(excl_amt),
                    "Tax Amount (₹)": float(tax_amt),
                    "Total Invoiced (₹)": float(tot_amt),
                    "Invoice Number": sanitize_formula_injection(inv.invoice_number),
                    "Invoice Date": str(inv.invoice_date),
                    "Due Date": str(inv.due_date),
                    "Cash Received (₹)": float(st["cash"]),
                    "TDS Deducted (₹)": float(st["tds"]),
                    "Total Settled (₹)": float(st["total"]),
                    "Balance Outstanding (₹)": float(bal),
                    "Status": status_str,
                    "Days Overdue": days_overdue,
                }
            )

        # 2. Unbilled Purchase Orders (included only in scope="all")
        if scope == "all":
            po_stmt = (
                select(PurchaseOrder, Customer, User)
                .join(Customer, PurchaseOrder.customer_id == Customer.id)
                .join(User, PurchaseOrder.executive_id == User.id)
                .where(PurchaseOrder.status == "UNBILLED")
                .order_by(PurchaseOrder.id.desc())
            )
            if fy and fy.upper() not in ("ALL", "ALL FY"):
                po_stmt = po_stmt.where(PurchaseOrder.fy_code == fy)
            if as_of_date:
                po_stmt = po_stmt.where(PurchaseOrder.po_date <= ref_date)

            po_rows = db.execute(po_stmt).all()

            for po, cust, exec_user in po_rows:
                records.append(
                    {
                        "FY": sanitize_formula_injection(po.fy_code),
                        "PO Date": str(po.po_date),
                        "PO Number": sanitize_formula_injection(po.po_number),
                        "Customer Name": sanitize_formula_injection(cust.billing_name),
                        "Executive": sanitize_formula_injection(exec_user.full_name or exec_user.username),
                        "Exclusive Amount (₹)": float(_quantize_2dec(po.exclusive_amount)),
                        "Tax Amount (₹)": float(_quantize_2dec(po.tax_amount)),
                        "Total Invoiced (₹)": 0.00,  # Unbilled PO has 0 recognized invoiced value
                        "Invoice Number": "—",
                        "Invoice Date": "—",
                        "Due Date": "—",
                        "Cash Received (₹)": 0.00,
                        "TDS Deducted (₹)": 0.00,
                        "Total Settled (₹)": 0.00,
                        "Balance Outstanding (₹)": 0.00,  # Unbilled PO has 0 recognized AR
                        "Status": "UNBILLED",
                        "Days Overdue": 0,
                    }
                )

    return records


@router.get("/excel")
def export_invoices_excel(
    request: Request,
    fy: str = "ALL",
    scope: str = "all",
    as_of_date: Optional[str] = None,
    user: dict = Depends(authenticate_and_consume_export_ticket),
):
    """
    Streams an authoritative 17-column Excel (.xlsx) workbook using openpyxl:
    - Single-use ticket consumed atomically.
    - Sanitized against formula injection.
    - Calibri 11pt, #1F4E79 header fill, and currency formatting.
    """
    rows = _compile_export_records(fy=fy, scope=scope, as_of_date=as_of_date)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CollectFlow Authoritative AR"

    h_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    h_fill = PatternFill(start_color=EXCEL_HEADER_FILL, end_color=EXCEL_HEADER_FILL, fill_type="solid")

    ws.append(AUTHORITATIVE_COLUMNS)
    for col_idx in range(1, len(AUTHORITATIVE_COLUMNS) + 1):
        cell = ws.cell(1, col_idx)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    currency_col_indices = [6, 7, 8, 12, 13, 14, 15]

    for r in rows:
        row_values = [
            sanitize_formula_injection(r["FY"]),
            r["PO Date"],
            sanitize_formula_injection(r["PO Number"]),
            sanitize_formula_injection(r["Customer Name"]),
            sanitize_formula_injection(r["Executive"]),
            r["Exclusive Amount (₹)"],
            r["Tax Amount (₹)"],
            r["Total Invoiced (₹)"],
            sanitize_formula_injection(r["Invoice Number"]),
            r["Invoice Date"],
            r["Due Date"],
            r["Cash Received (₹)"],
            r["TDS Deducted (₹)"],
            r["Total Settled (₹)"],
            r["Balance Outstanding (₹)"],
            sanitize_formula_injection(r["Status"]),
            r["Days Overdue"],
        ]
        ws.append(row_values)

        current_row = ws.max_row
        for c_idx in currency_col_indices:
            cell = ws.cell(current_row, c_idx)
            cell.number_format = EXCEL_CURRENCY_FORMAT

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"CollectFlow-Authoritative-AR-{fy.replace(' ', '_')}-{datetime.date.today().isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/csv")
def export_invoices_csv(
    request: Request,
    fy: str = "ALL",
    scope: str = "all",
    as_of_date: Optional[str] = None,
    user: dict = Depends(authenticate_and_consume_export_ticket),
):
    """
    Streams an authoritative 17-column CSV:
    - Same records and ordering as Excel export.
    - Sanitized against spreadsheet formula injection.
    - Single-use ticket consumed atomically.
    """
    rows = _compile_export_records(fy=fy, scope=scope, as_of_date=as_of_date)

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(AUTHORITATIVE_COLUMNS)

    for r in rows:
        writer.writerow(
            [
                sanitize_formula_injection(r["FY"]),
                r["PO Date"],
                sanitize_formula_injection(r["PO Number"]),
                sanitize_formula_injection(r["Customer Name"]),
                sanitize_formula_injection(r["Executive"]),
                _format_money(r["Exclusive Amount (₹)"]),
                _format_money(r["Tax Amount (₹)"]),
                _format_money(r["Total Invoiced (₹)"]),
                sanitize_formula_injection(r["Invoice Number"]),
                r["Invoice Date"],
                r["Due Date"],
                _format_money(r["Cash Received (₹)"]),
                _format_money(r["TDS Deducted (₹)"]),
                _format_money(r["Total Settled (₹)"]),
                _format_money(r["Balance Outstanding (₹)"]),
                sanitize_formula_injection(r["Status"]),
                r["Days Overdue"],
            ]
        )

    output.seek(0)
    filename = f"CollectFlow-Authoritative-AR-{fy.replace(' ', '_')}-{datetime.date.today().isoformat()}.csv"
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
