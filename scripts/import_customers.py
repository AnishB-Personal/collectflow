#!/usr/bin/env python3
"""
Import customer records from the original Excel workbook into CollectFlow database.
Preserves existing records, avoids duplicates by lowercase billing_name,
and populates complete contact and tax attributes.
"""

import os
import re
import sqlite3
import sys
from pathlib import Path
import openpyxl


def clean_str(val):
    if val is None:
        return None
    s = str(val).strip()
    return s if s else None


def clean_gst(val):
    s = clean_str(val)
    if not s:
        return None
    s = s.upper()
    if s in ("NONE", "NIL", "NA", "-", "N/A", "0"):
        return None
    return s


def clean_pin(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return str(int(val))
    s = str(val).strip()
    return s if s else None


def clean_phone(val):
    if val is None:
        return None
    if isinstance(val, float):
        return str(int(val))
    s = str(val).strip()
    # Normalize excessive spaces
    s = re.sub(r"\s+", " ", s)
    return s if s else None


def import_customers(excel_path: str, db_path: str):
    print(f"Loading Excel: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if "Customers" not in wb.sheetnames:
        print("ERROR: 'Customers' sheet not found in workbook.")
        sys.exit(1)

    ws = wb["Customers"]

    print(f"Connecting to DB: {db_path}")
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    # Ensure email column exists
    cols = [r["name"] for r in cur.execute("PRAGMA table_info(customers);").fetchall()]
    if "email" not in cols:
        print("Adding 'email' column to customers table...")
        cur.execute("ALTER TABLE customers ADD COLUMN email TEXT;")
        conn.commit()

    # Load existing customer names from DB
    existing_rows = cur.execute("SELECT id, LOWER(TRIM(billing_name)) as norm_name FROM customers;").fetchall()
    existing_map = {r["norm_name"]: r["id"] for r in existing_rows}
    print(f"Existing customers in DB: {len(existing_map)}")

    inserted_count = 0
    updated_count = 0
    skipped_count = 0

    # Row 4 is header, data starts at row 5
    for row_idx, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
        raw_name = row[1]
        name = clean_str(raw_name)
        if not name:
            skipped_count += 1
            continue

        addr = clean_str(row[2])
        city = clean_str(row[3])
        state = clean_str(row[4])
        country = clean_str(row[5]) or "IN"
        pin = clean_pin(row[6])
        gst = clean_gst(row[7])

        phone_01 = clean_phone(row[8])
        phone_02 = clean_phone(row[9])
        contact_name = clean_str(row[10])
        designation = clean_str(row[11])
        mobile = clean_phone(row[12])
        email_id = clean_str(row[13])

        # Secondary contact fields if primary is empty
        sec_contact = clean_str(row[14]) if len(row) > 14 else None
        sec_desig = clean_str(row[15]) if len(row) > 15 else None
        sec_mobile = clean_phone(row[16]) if len(row) > 16 else None
        sec_email = clean_str(row[17]) if len(row) > 17 else None

        # Phone consolidation
        p1 = phone_01 if phone_01 else mobile
        p2 = phone_02 if phone_02 else (mobile if phone_01 and phone_01 != mobile else sec_mobile)
        email = email_id if email_id else sec_email
        contact = contact_name if contact_name else sec_contact
        desig = designation if designation else sec_desig

        norm_name = name.lower()

        if norm_name in existing_map:
            cid = existing_map[norm_name]
            # Update customer with richer information if present
            cur.execute("""
                UPDATE customers SET
                    billing_address = COALESCE(billing_address, ?),
                    city = COALESCE(city, ?),
                    state = COALESCE(state, ?),
                    country = COALESCE(country, ?),
                    pin_code = COALESCE(pin_code, ?),
                    gst_number = COALESCE(gst_number, ?),
                    phone_1 = COALESCE(phone_1, ?),
                    phone_2 = COALESCE(phone_2, ?),
                    contact_person = COALESCE(contact_person, ?),
                    designation = COALESCE(designation, ?),
                    email = COALESCE(email, ?)
                WHERE id = ?;
            """, (addr, city, state, country, pin, gst, p1, p2, contact, desig, email, cid))
            updated_count += 1
        else:
            cur.execute("""
                INSERT INTO customers (
                    billing_name, billing_address, city, state, country,
                    pin_code, gst_number, phone_1, phone_2, contact_person,
                    designation, email
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
            """, (name, addr, city, state, country, pin, gst, p1, p2, contact, desig, email))
            new_id = cur.lastrowid
            existing_map[norm_name] = new_id
            inserted_count += 1

    conn.commit()
    conn.close()

    print(f"=== Import Completed ===")
    print(f"Inserted: {inserted_count}")
    print(f"Updated:  {updated_count}")
    print(f"Skipped:  {skipped_count}")
    print(f"Total customers in DB now: {len(existing_map)}")


if __name__ == "__main__":
    base_dir = Path(__file__).resolve().parent.parent
    default_excel = base_dir.parent / "PO_Payment_Details_FY18-19_to_26-27_Upgraded.xlsx"
    default_db = base_dir / "data" / "collectflow.db"

    excel_file = sys.argv[1] if len(sys.argv) > 1 else str(default_excel)
    db_file = sys.argv[2] if len(sys.argv) > 2 else str(default_db)

    import_customers(excel_file, db_file)
